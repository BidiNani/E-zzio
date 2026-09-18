"""
E-ZZIO Sovereign Generator — Excel (XLSX) Spreadsheet Engine.
Génère des classeurs Excel professionnels, stylisés et avec formules via openpyxl.
"""
from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Any

logger = logging.getLogger("SheetEngine")


class SheetEngine:
    """Moteur souverain de génération de classeurs XLSX."""

    def __init__(self, workspace_root: str = "G:\\AI\\E-zzio"):
        self.workspace_root = Path(workspace_root).resolve()
        self.exports_dir = self.workspace_root / "outputs"
        self.exports_dir.mkdir(parents=True, exist_ok=True)

    def is_available(self) -> bool:
        """Vérifie si la bibliothèque openpyxl est installée."""
        try:
            import openpyxl  # noqa: F401
            return True
        except ImportError:
            return False

    def generate_spreadsheet(
        self,
        filename: str,
        sheets_data: list[dict[str, Any]],
        title: str | None = None
    ) -> dict[str, Any]:
        """
        Génère un classeur XLSX multi-onglets avec mise en forme professionnelle.
        sheets_data: [
            {
                "sheet_name": "Revenus",
                "headers": ["Mois", "Ventes", "Dépenses", "Bénéfice"],
                "rows": [["Janvier", 12000, 8000, "=B2-C2"], ["Février", 15000, 9000, "=B3-C3"]],
                "totals": ["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)"]
            }
        ]
        """
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as e:
            raise RuntimeError(
                "[SHEET-ENGINE] Dépendance 'openpyxl' manquante. "
                "Installez-la via `pip install openpyxl`."
            ) from e

        clean_name = os.path.basename(filename.strip())
        if not clean_name.endswith(".xlsx"):
            clean_name += ".xlsx"

        out_path = self.exports_dir / clean_name

        wb = openpyxl.Workbook()
        # Supprime la feuille par défaut vide
        wb.remove(wb.active)

        # Palettes de styles Cyberpunk / Pro E-ZzIO
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="38BDF8")
        total_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        total_font = Font(name="Segoe UI", size=11, bold=True, color="34D399")
        cell_font = Font(name="Segoe UI", size=10, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for s_idx, s_info in enumerate(sheets_data):
            sheet_name = s_info.get("sheet_name", f"Feuille_{s_idx+1}")[:31]
            ws = wb.create_sheet(title=sheet_name)
            ws.views.sheetView[0].showGridLines = True

            current_row = 1
            # Titre optionnel
            if title and s_idx == 0:
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(s_info.get("headers", [])), 4))
                title_cell = ws.cell(row=1, column=1, value=title)
                title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0EA5E9")
                title_cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[1].height = 30
                current_row = 3

            # En-têtes
            headers = s_info.get("headers", [])
            if headers:
                for col_idx, h in enumerate(headers, start=1):
                    c = ws.cell(row=current_row, column=col_idx, value=h)
                    c.fill = header_fill
                    c.font = header_font
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border = thin_border
                ws.row_dimensions[current_row].height = 24
                current_row += 1

            # Lignes de données
            rows = s_info.get("rows", [])
            for r_data in rows:
                for col_idx, val in enumerate(r_data, start=1):
                    c = ws.cell(row=current_row, column=col_idx, value=val)
                    c.font = cell_font
                    c.border = thin_border
                    if isinstance(val, (int, float)):
                        c.alignment = Alignment(horizontal="right")
                    else:
                        c.alignment = Alignment(horizontal="left")
                ws.row_dimensions[current_row].height = 19
                current_row += 1

            # Ligne de totaux optionnelle
            totals = s_info.get("totals", [])
            if totals:
                for col_idx, t_val in enumerate(totals, start=1):
                    c = ws.cell(row=current_row, column=col_idx, value=t_val)
                    c.fill = total_fill
                    c.font = total_font
                    c.border = thin_border
                ws.row_dimensions[current_row].height = 22

            # Ajustement automatique des largeurs de colonnes
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(out_path))
        logger.info("[SHEET-ENGINE] Classeur généré avec succès : %s", out_path)

        return {
            "ok": True,
            "filename": clean_name,
            "path": str(out_path),
            "size_bytes": out_path.stat().st_size,
            "sheets": [s.title for s in wb.worksheets]
        }
